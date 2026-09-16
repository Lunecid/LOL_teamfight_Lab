"""Build the collaborator DOCX from the audited English Markdown source."""
from pathlib import Path
import hashlib
import json
import re
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/COLLABORATOR_RESEARCH_SPECIFICATION_20260915.md'
OUT = ROOT / 'deliverables/collaborator_specification_20260915'
QA = ROOT / 'outputs/collaborator_specification_20260915'
OUT.mkdir(parents=True, exist_ok=True)
QA.mkdir(parents=True, exist_ok=True)
text = SOURCE.read_text(encoding='utf-8')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.68)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.3)
styles = doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    styles[name].font.name = 'Calibri'
    styles[name].font.color.rgb = RGBColor(0,0,0)
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.line_spacing = 1.08
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.widow_control = True
styles['Title'].font.size = Pt(25)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(10)
for border in list(styles['Title'].element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
styles['Subtitle'].font.size = Pt(12)
for name,size in [('Heading 1',16),('Heading 2',12.5),('Heading 3',11.5)]:
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].paragraph_format.space_before = Pt(13)
    styles[name].paragraph_format.space_after = Pt(6)
    styles[name].paragraph_format.keep_with_next = True
header = sec.header.paragraphs[0]
header.text = 'ENGAGEMENT PREDICTION  |  RESEARCH SPECIFICATION'
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('15 September 2026  |  Page ').font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
footer._p.append(field)
doc.core_properties.title = 'League of Legends Engagement Prediction Research Specification'
doc.core_properties.author = 'Seongeun'
doc.core_properties.subject = 'Implemented methods and results with separately identified planned experiments'
doc.core_properties.keywords = 'League of Legends, engagement, win probability, labels, SHAP'

def hyperlink(p,label,url):
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),'222222'); pr.append(c)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); pr.append(u)
    r.append(pr); t=OxmlElement('w:t'); t.text=label; r.append(t);h.append(r);p._p.append(h)

def inline(p,s):
    tokens=re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\))',s)
    for part in tokens:
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if m: hyperlink(p,m[1],m[2])
        elif part.startswith('**') and part.endswith('**'): p.add_run(part[2:-2]).bold=True
        else: p.add_run(part)

def table(rows):
    cols=len(rows[0]); tab=doc.add_table(rows=0,cols=cols)
    tab.style='Table Grid'; tab.autofit=False
    if cols==6: widths=[1.55,1.06,1.12,1.15,1.1,1.12]
    elif cols==5: widths=[2.14,1.24,1.24,1.24,1.24]
    elif cols==4: widths=[1.5,1.35,1.4,2.85]
    elif cols==3: widths=[2.1,2.15,2.85]
    else: widths=[4.9,2.2]
    for col,w in zip(tab.columns,widths): col.width=Inches(w)
    for ri,row in enumerate(rows):
        cells=tab.add_row().cells
        trpr=tab.rows[-1]._tr.get_or_add_trPr()
        if ri==0:
            rep=OxmlElement('w:tblHeader');trpr.append(rep)
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        for ci,val in enumerate(row):
            cells[ci].width=Inches(widths[ci])
            p=cells[ci].paragraphs[0]
            p.paragraph_format.space_after=Pt(3)
            p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.line_spacing=1.02
            inline(p,val)
            for r in p.runs:
                r.font.size=Pt(9.5 if cols>=5 else 10)
                r.font.bold=(ri==0)
            tcpr=cells[ci]._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');tcpr.append(shade)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def mr(value):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;r.append(t);return r

def sub(base,index):
    n=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(base));s=OxmlElement('m:sub');s.append(mr(index));n.extend([e,s]);return n

def frac(top,bottom):
    n=OxmlElement('m:f');a=OxmlElement('m:num');a.extend(top);b=OxmlElement('m:den');b.extend(bottom);n.extend([a,b]);return n

def mathline(items):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together=True
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(6)
    om=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath')
    m.extend([mr(x) if isinstance(x,str) else x for x in items]);om.append(m);p._p.append(om)

def equation(number):
    if number==1:
        mathline([sub('Y','CoG,i'),' = 1[ ',sub('Σ','u∈i'),sub('α','u'),sub('σ','u'),sub('v','u'),' > 0 ]'])
        mathline([sub('α','u'),' = ',frac([mr('exp(2'),sub('p','u'),mr(')')],[sub('Σ','v∈i'),mr('exp(2'),sub('p','v'),mr(')')])])
    elif number==2:
        mathline([sub('p','t'),' = ',sub('V','θ'),'(S(t)) ≈ P(W = 1 | S(t))'])
        mathline([sub('q','ϕ'),'(',sub('X','pre,i'),') ≈ P(',sub('Y','i'),' = 1 | ',sub('X','pre,i'),')'])
    elif number==3:
        mathline([sub('t','pre,i'),' = ',sub('K','i'),' − 15 s − 1 ms'])
        mathline([sub('e','i'),'(h) = min{',sub('L','i'),' + h, ',sub('J','i'),' − 1 ms, ',sub('S','next,i'),' − 1 ms, ',sub('T','end,i'),' − 1 ms}'])
    elif number==4:
        mathline([sub('ΔV','i'),'(h) = ',sub('V','θᵢ'),'(S(',sub('e','i'),'(h))) − ',sub('V','θᵢ'),'(S(',sub('t','pre,i'),'))'])
        mathline([sub('Y','i'),'(h) = 1[',sub('ΔV','i'),'(h) > 0]'])
    elif number==5:
        mathline([sub('ℒ','q'),' = −',frac([mr('1')],[sub('Σ','i'),sub('w','i')]),' ',sub('Σ','i'),sub('w','i'),' {',sub('Y','i'),' log ',sub('q','i'),' + (1 − ',sub('Y','i'),') log(1 − ',sub('q','i'),')}'])
    else: raise RuntimeError('Unspecified equation')

eq=0;lines=text.splitlines();i=0
while i<len(lines):
    s=lines[i].strip()
    if not s: i+=1;continue
    if s.startswith('$$'):
        eq+=1
        equation(eq)
        i+=1;continue
    if s.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',c) for c in cells):rows.append(cells)
            i+=1
        table(rows);continue
    if s.startswith('# '):
        p=doc.add_paragraph(s[2:],style='Title')
    elif s.startswith('## '):
        p=doc.add_paragraph(s[3:],style='Heading 1')
    elif s.startswith('### '):
        p=doc.add_paragraph(s[4:],style='Heading 2')
    else:
        p=doc.add_paragraph();inline(p,s)
    i+=1

docx=OUT/'Research_Specification_20260915.docx'
doc.save(docx)
shutil.copyfile(SOURCE,OUT/'Research_Specification_20260915.md')
schema=ROOT/'outputs/full_corpus_training_20260915/q_pre_only_schema.json'
x=json.loads(schema.read_text(encoding='utf-8'))
public={k:x[k] for k in ['input_names_all','predictor_sets','predictor_sets_sha256','ridge_count','economic_count','categorical_excluded','feature_units','shap_groups','excluded'] if k in x}
public['scope']='Ordered q feature dictionary from the completed full-corpus run. Names and definitions only; no match records. input_names_all is the q assembly schema, not a list of all raw state channels.'
public['source_sha256']=hashlib.sha256(schema.read_bytes()).hexdigest()
v_source=ROOT/'outputs/full_corpus_training_20260915/v_models_manifest.json'
v=json.loads(v_source.read_text(encoding='utf-8'))
assert len(v['feature_names'])==361
public['V_input_names']=v['feature_names']
public['V_input_count']=361
public['V_source_sha256']=hashlib.sha256(v_source.read_bytes()).hexdigest()
public['StateV2_audit_only_column']='snapshot_age_s'
public['StateV2_column_count']=362
(OUT/'Feature_Dictionary_20260915.json').write_text(json.dumps(public,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
sources=[SOURCE,ROOT/'docs/TOG_END_TO_END_READINESS_AUDIT_20260915.md',ROOT/'docs/LABEL_VALIDITY_FINDINGS_20260915.md',ROOT/'docs/OBJECTIVE_CHANNEL_ABLATION_FINDINGS_20260915.md',ROOT/'outputs/cohort_role_training_20260915/REPORT.md',ROOT/'outputs/objective_channel_ablation_20260915/validation.json',ROOT/'tmp/pdfs/paper_308_review/paper_308_text.txt',ROOT/'docs/DELTA_Q_FAIR_BALANCED_PROTOCOL_20260915.md',ROOT/'docs/COG_MODEL_COVERAGE_DELTA_Q_20260915.md',schema,v_source]
manifest={'created':'2026-09-15','purpose':'Collaborator specification source audit','sources':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},'outputs':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.iterdir() if p.is_file()},'review':{'main_results':'Luna bounded evidence check and root source review','human_semantic_review':False,'new_experiments_launched':False,'layout':'pending'}}
(QA/'source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx':str(docx),'words':len(text.split()),'equations':eq,'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)}))
